import json
import os
import argparse
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import (PatternFill, Font, Alignment,
                                  Border, Side)
except ImportError:
    raise SystemExit("Run: pip install openpyxl")

OUTPUT_DIR = Path(__file__).parent.parent / "output"

# ── Colours ──────────────────────────────────────────────────────────────────
COLOURS = {
    "header_bg":   "1F3864",   # dark navy
    "header_fg":   "FFFFFF",
    "positive":    "C6EFCE",   # green
    "negative":    "FECDCD",   # red
    "edge_case":   "FFEB9C",   # yellow
    "high":        "FF0000",
    "medium":      "FF9900",
    "low":         "70AD47",
    "alt_row":     "F2F2F2",
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def col_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def export_to_excel(json_path: str, output_path: str = None):
    with open(json_path, "r") as f:
        data = json.load(f)

    feature    = data.get("feature", "Test Cases")
    test_cases = data.get("test_cases", [])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    # ── Title row ────────────────────────────────────────────────────────────
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value         = f"AI Generated Test Cases — {feature}"
    title_cell.font          = Font(bold=True, size=13, color="FFFFFF")
    title_cell.fill          = col_fill("1F3864")
    title_cell.alignment     = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # ── Sub-header: generated date ────────────────────────────────────────────
    ws.merge_cells("A2:H2")
    date_cell = ws["A2"]
    date_cell.value      = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  Total: {len(test_cases)} test cases"
    date_cell.font       = Font(italic=True, size=9, color="595959")
    date_cell.fill       = col_fill("DCE6F1")
    date_cell.alignment  = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 16

    # ── Column headers ────────────────────────────────────────────────────────
    headers = ["ID", "Title", "Type", "Priority",
               "Preconditions", "Steps", "Expected Result", "Tags"]
    col_widths = [8, 35, 12, 10, 25, 40, 30, 20]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font      = Font(bold=True, color=COLOURS["header_fg"], size=10)
        cell.fill      = col_fill(COLOURS["header_bg"])
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = BORDER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[3].height = 20

    # ── Data rows ─────────────────────────────────────────────────────────────
    for row_idx, tc in enumerate(test_cases, start=4):
        is_alt = (row_idx % 2 == 0)
        row_bg = COLOURS["alt_row"] if is_alt else "FFFFFF"

        steps_text = "\n".join(
            [f"{i+1}. {s}" for i, s in enumerate(tc.get("steps", []))]
        )
        tags_text = ", ".join(tc.get("tags", []))

        row_data = [
            tc.get("id", ""),
            tc.get("title", ""),
            tc.get("type", ""),
            tc.get("priority", ""),
            tc.get("preconditions", ""),
            steps_text,
            tc.get("expected_result", ""),
            tags_text,
        ]

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border    = BORDER

            # Default row background
            cell.fill = col_fill(row_bg)

            # Type column colour
            if col_idx == 3:
                type_color = COLOURS.get(tc.get("type", ""), row_bg)
                cell.fill      = col_fill(type_color)
                cell.alignment = Alignment(horizontal="center", vertical="top")

            # Priority column colour
            if col_idx == 4:
                pri_color = COLOURS.get(tc.get("priority", "").lower(), row_bg)
                cell.font      = Font(bold=True, color=pri_color)
                cell.alignment = Alignment(horizontal="center", vertical="top")

            # ID column bold
            if col_idx == 1:
                cell.font      = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="top")

        # Row height based on step count
        step_count = len(tc.get("steps", []))
        ws.row_dimensions[row_idx].height = max(40, step_count * 15)

    # ── Freeze header rows ────────────────────────────────────────────────────
    ws.freeze_panes = "A4"

    # ── Auto filter ───────────────────────────────────────────────────────────
    ws.auto_filter.ref = f"A3:H{3 + len(test_cases)}"

    # ── Save ──────────────────────────────────────────────────────────────────
    if not output_path:
        timestamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name   = feature[:30].replace(" ", "_").lower()
        output_path = str(OUTPUT_DIR / f"testcases_{safe_name}_{timestamp}.xlsx")

    OUTPUT_DIR.mkdir(exist_ok=True)
    wb.save(output_path)
    print(f"\nExcel saved to: {output_path}")
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Export test cases JSON to Excel")
    parser.add_argument("--json", type=str, required=True,
                        help="Path to the JSON file in output/ folder")
    args = parser.parse_args()
    export_to_excel(args.json)


if __name__ == "__main__":
    main()